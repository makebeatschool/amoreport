import os
import time

import requests
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment
from openpyxl.utils import get_column_letter
from amoCRMWork import load_events_for_user
from calculate import calculate_lead_stats, calculate_stats_for_manager, calculate_stats_for_telemarketer
from config import TELEGRAM_DOMEN, PARENT_DIR
from env import TELEGRAM_BOT_TOKEN, TELEGRAM_CHAT_ID
from data import MANAGERS, TELEMARKETERS

BASE_DATA = {
    "Менеджеры":{},
    "Телемаркетологи":{}
}

def send_telegram_file(filepath):
    url = f"{TELEGRAM_DOMEN}/bot{TELEGRAM_BOT_TOKEN}/sendDocument"
    with open(filepath, "rb") as f:
        requests.post( url,
            data={"chat_id": TELEGRAM_CHAT_ID},
            files={"document": f}
        )

def build_report_excel(data, date_range):
    wb = Workbook()
    gray_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    for sheet_name, employees in data.items():
        ws = wb.active if sheet_name == list(data.keys())[0] else wb.create_sheet()
        ws.title = sheet_name
        employees_list = list(employees.keys())
        all_fields = set()
        for emp in employees.values():
            all_fields.update(emp.keys())
        all_fields = list(all_fields)
        ws.cell(row=1, column=1, value="")
        for col, emp in enumerate(employees_list, start=2):
            ws.cell(row=1, column=col, value=emp)
        for row, field in enumerate(all_fields, start=2):
            ws.cell(row=row, column=1, value=field)
            for col, emp in enumerate(employees_list, start=2):
                value = employees.get(emp, {}).get(field, "")
                cell = ws.cell(row=row, column=col, value=value)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if col % 2 == 0:
                    cell.fill = gray_fill
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max_len + 2
        for row in ws.iter_rows():
            ws.row_dimensions[row[0].row].height = 18
        ws.freeze_panes = "B2"
    filename = os.path.join(PARENT_DIR, '', f'ОТЧЁТ_ЗА_{date_range}.xlsx')
    wb.save(filename)
    return filename

def format_manager_report_exel(stats, name, baseData):
    baseData["Менеджеры"][name] = {
        "Звонков": stats['total_calls'],
        "Исходящие": stats['outgoing'],
        "Входящие": stats['incoming'],
        "Диалоги": stats['answered'],
        "Диалоги от 25 сек": stats['dialogs_25s'],
        "Уникальные диалоги": len(stats['unique_contacts']),
        "Недозвон": stats['missed'],
        "ОМА": stats['oma'],
        "СОМА": stats['soma'],
        "СОГА": stats['soga'],
        "ОГА": stats['oga'],
        "Отказы": stats['refusals'],
        "Отказы после 1 оплаты": stats['cancel_after_first_payment'],
        "Думают": stats['thinking'],
        "Новые лиды звонки": stats['new_leads']['calls'],
        "Новые лиды диалоги": stats['new_leads']['dialogs'],
        "Старые лиды звонки": stats['old_leads']['calls'],
        "Старые лиды диалоги": stats['old_leads']['dialogs']
    }
    return baseData

def format_telemarketer_report_exel(stats, name, baseData):
    baseData["Телемаркетологи"][name] = {
        "Звонков": stats['total_calls'],
        "Исходящие": stats['outgoing'],
        "Входящие": stats['incoming'],
        "Диалоги": stats['answered'],
        "Диалоги от 25 сек": stats['dialogs_25s'],
        "Уникальные диалоги": len(stats['unique_contacts']),
        "Недозвон": stats['missed'],

        "Отказы": stats['refusals'],
        "Думают": stats['thinking'],
        "Согласие на пробное": stats['soglasiy_na_probnoe'],
        "Дошло до пробника": stats['doshlo_do_probnika'],
        "Перезапись 1": stats['perezapis_1'],
        "Перезапись 2": stats['perezapis_2'],

        "Новые лиды звонки": stats['new_leads']['calls'],
        "Новые лиды диалоги": stats['new_leads']['dialogs'],
        "Новые лиды согласия": stats['new_leads']['soglasiye'],
        "Новые лиды дошло до пробника": stats['new_leads']['doshlo'],

        "Старые лиды звонки": stats['old_leads']['calls'],
        "Старые лиды диалоги": stats['old_leads']['dialogs'],
        "Старые лиды согласия": stats['old_leads']['soglasiye'],
        "Старые лиды дошло до пробника": stats['old_leads']['doshlo'],

        "Добавлены в группу": len(stats['added_group_contacts']),

        "Конверсия в согласие": stats['new_leads']['conversion_to_soglasiye'],
        "Конверсия в дозвон": stats['new_leads']['conversion_to_dialogs'],
        "Конверсия диалогов": stats['new_leads']['conversion_from_dialogs']
    }
    return baseData

def send_report_exel(date_from, date_to, date_range, employers, calul_func, format_func, basedata):
    for emp in employers:
        user_id = emp['id']
        user_name = emp['name']
        print(f"  📌 {user_name}...")
        events = load_events_for_user(user_id, date_from, date_to, max_pages=10)
        if not events:
            print(f"    ⚠️ Нет событий, пропускаем")
            continue
        contact_stats = calculate_lead_stats(events)
        stats = calul_func(events, contact_stats)
        basedata = format_func(stats, user_name, basedata)
        time.sleep(1)
    return basedata

def send_telegram_exel(date_from, date_to, date_range):
    print("Генерация эксель файла")
    basedata = send_report_exel(date_from, date_to, date_range, MANAGERS, 
                calculate_stats_for_manager, format_manager_report_exel, BASE_DATA)
    basedata = send_report_exel(date_from, date_to, date_range, TELEMARKETERS, 
                 calculate_stats_for_telemarketer, format_telemarketer_report_exel, basedata)
    filename = build_report_excel(basedata, date_range)
    print("Отправка эксель файла")
    send_telegram_file(filename)
    if os.path.exists(filename):
        os.remove(filename)